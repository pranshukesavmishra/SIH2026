"""National-round deck: official SIH template + the team's approved graphics
(cropped at 300 dpi from their 10-Sept deck) + every content fix from
docs/PROJECT_STATE.md section 4. Single source of truth — rebuilds from
template.pptx each run."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from urllib.parse import quote_plus

NAVY  = RGBColor(0x0B, 0x25, 0x45)
CYAN  = RGBColor(0x0E, 0x7C, 0x99)
GREEN = RGBColor(0x1F, 0x8A, 0x5B)
GRAY  = RGBColor(0x5A, 0x64, 0x72)
LIGHT = RGBColor(0xEF, 0xF3, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED   = RGBColor(0x9C, 0x2E, 0x23)

prs = Presentation("template.pptx")
S = list(prs.slides)


def set_runs(par, runs, size=14, align=None):
    for r in list(par.runs):
        r._r.getparent().remove(r._r)
    if align is not None:
        par.alignment = align
    for text, opt in runs:
        r = par.add_run(); r.text = text
        f = r.font
        f.name = "Arial"
        f.size = Pt(opt.get("size", size))
        f.bold = opt.get("bold", False)
        f.italic = opt.get("italic", False)
        f.color.rgb = opt.get("color", NAVY)


def tb(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    return box, tf


def card(slide, x, y, w, h, fill=LIGHT, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = 0.06
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(1)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = Inches(0.12)
    tf.margin_top = tf.margin_bottom = Inches(0.08)
    return shp, tf


def bullets(tf, items, size=13, gap=4, color=NAVY):
    first = True
    for item in items:
        p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
        first = False
        p.space_after = Pt(gap)
        if isinstance(item, tuple):
            lead, rest = item
            set_runs(p, [(lead, {"bold": True, "size": size, "color": color}),
                         (rest, {"size": size, "color": NAVY})])
        else:
            set_runs(p, [(item, {"size": size, "color": NAVY})])


def find(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None


def style_pointer_box(slide, y=1.28, h=0.40, size=9):
    box = find(slide, "TextBox 8")
    box.left, box.top = Inches(1.92), Inches(max(y, 1.28))
    box.width, box.height = Inches(11.0), Inches(h)
    tf = box.text_frame
    tf.word_wrap = True
    for p in tf.paragraphs:
        text = "".join(r.text for r in p.runs)
        if not text.strip():
            p._p.getparent().remove(p._p)
            continue
        set_runs(p, [(text.strip(), {"italic": True, "size": size, "color": GRAY})])
        p.space_after = Pt(0)


def team_oval(slide):
    oval = None
    for sh in slide.shapes:
        if sh.name.startswith("Oval"):
            oval = sh
    if oval is not None:
        oval.fill.solid(); oval.fill.fore_color.rgb = NAVY
        oval.line.color.rgb = CYAN; oval.line.width = Pt(1.25)
        tf = oval.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for extra in list(tf.paragraphs[1:]):
            extra._p.getparent().remove(extra._p)
        set_runs(tf.paragraphs[0], [("ZeroDrift", {"size": 12, "bold": True, "color": WHITE})],
                 align=PP_ALIGN.CENTER)


def clean_title(slide):
    t = find(slide, "Title 1")
    if t is not None:
        for par in t.text_frame.paragraphs:
            for r in par.runs:
                r.text = r.text.replace("\x0b", "").replace("", "")


def pic(slide, path, x, y, w=None, h=None, border=None):
    kw = {}
    if w is not None:
        kw["width"] = Inches(w)
    if h is not None:
        kw["height"] = Inches(h)
    p = slide.shapes.add_picture(path, Inches(x), Inches(y), **kw)
    if border is not None:
        p.line.color.rgb = border; p.line.width = Pt(0.75)
    return p


# ================= SLIDE 1 — TITLE =================
s = S[0]
title = find(s, "Title 7")
for par in title.text_frame.paragraphs:
    for r in par.runs:
        r.font.size = Pt(34)
title.width = Inches(10.2)
FIELDS = {
    "Problem Statement ID": "  26169",
    "Problem Statement Title": ("  AI-Based Virtual Camera Tracking System for Coarse "
                                "Alignment of Mobile FSOC Terminals"),
    "Theme": "  Smart Automation",
    "PS Category": "  Software",
    "Team ID": "  —",
    "Team Name": "  ZeroDrift",
}
box = find(s, "TextBox 9")
for p in box.text_frame.paragraphs:
    text = "".join(r.text for r in p.runs).strip()
    matched = None
    for key, val in FIELDS.items():
        if text.startswith(key):
            matched = (key, val); break
    if matched is None:
        continue
    key, val = matched
    set_runs(p, [(key + " –", {"bold": True, "size": 14, "color": NAVY}),
                 (val, {"size": 14, "color": CYAN,
                        "bold": key in ("Team Name", "Problem Statement ID")})])
    p.space_after = Pt(6)
    p.line_spacing = 1.0

shp, tf = card(s, 0.36, 5.62, 5.85, 1.10, fill=LIGHT)
set_runs(tf.paragraphs[0], [("The problem in one line:  ", {"bold": True, "size": 11.5, "color": NAVY}),
                            ("a mobile laser (FSOC) terminal must find, verify and continuously "
                             "point at its partner's beacon with hair-thin accuracy — ISRO asks "
                             "for this coarse-alignment brain, proven entirely in software.",
                             {"size": 11.5, "color": NAVY})])
pic(s, "c1_isro.jpg", 6.55, 6.32, h=0.95)
pic(s, "logo_banner.png", 8.35, 6.28, h=1.02)

# ================= SLIDE 2 — IDEA =================
s = S[1]
clean_title(s); team_oval(s)
style_pointer_box(s, h=0.60, size=9)

_, tf = tb(s, 0.30, 1.94, 6.75, 0.44)
set_runs(tf.paragraphs[0],
         [("Problem in one line:  ", {"bold": True, "size": 11, "color": RED}),
          ("find the partner's blinking beacon among stars and glints, verify it, and hold "
           "it centred to micro-radian accuracy — automatically, in software.",
           {"size": 11, "color": NAVY})])
_, tf = tb(s, 0.30, 2.44, 6.75, 0.44)
set_runs(tf.paragraphs[0],
         [("Our answer — the eyes & neck of a laser terminal:  ",
           {"bold": True, "size": 12, "color": CYAN}),
          ("simulate the sky · verify the beacon · track & predict.  Built, measured, live.",
           {"size": 11.5, "color": NAVY})])

pic(s, "c2_cards.jpg", 0.32, 2.94, w=6.30)                       # h = 3.16
for i, chp in enumerate(["c2_chip1p.jpg", "c2_chip2p.jpg", "c2_chip3.jpg"]):
    pic(s, chp, 0.32 + i * 2.19, 6.16, w=2.05)

pic(s, "c2_shot.jpg", 7.42, 2.02, h=3.00, border=GRAY)           # w = 4.90
_, tf = tb(s, 7.05, 5.08, 5.90, 0.26)
set_runs(tf.paragraphs[0], [("Our console locked on the beacon — the exact view the live demo link opens",
                             {"italic": True, "size": 10, "color": GRAY})], align=PP_ALIGN.CENTER)
_, tf = tb(s, 7.05, 5.36, 5.90, 0.24)
set_runs(tf.paragraphs[0],
         [("Measured: ", {"bold": True, "size": 11, "color": NAVY}),
          ("median 198 µrad", {"bold": True, "size": 11, "color": CYAN}),
          (" in lock · in-FOV ", {"size": 11, "color": NAVY}),
          ("100%", {"bold": True, "size": 11, "color": GREEN}),
          (" · the live-demo run", {"size": 11, "color": NAVY})], align=PP_ALIGN.CENTER)
pic(s, "chart_convergence.png", 7.05, 5.64, w=5.90,
    border=RGBColor(0xD5, 0xDD, 0xE5))                            # h = 1.30

# ================= SLIDE 3 — TECHNICAL =================
s = S[2]
clean_title(s); team_oval(s)
style_pointer_box(s, h=0.26, size=9)
pic(s, "c3_all.jpg", 1.20, 1.72, w=10.94)                         # h = 5.16

# ================= SLIDE 4 — FEASIBILITY =================
s = S[3]
clean_title(s); team_oval(s)
style_pointer_box(s, h=0.26, size=9)
pic(s, "c4_cols.jpg", 1.05, 1.62, w=11.22)                        # h = 4.32
pic(s, "c4_strip.jpg", 3.42, 6.00, w=6.50)                        # h = 0.93
_, tf = tb(s, 0.35, 6.06, 2.95, 0.80, anchor=MSO_ANCHOR.MIDDLE)
bullets(tf, [("Worst case, 64 runs →  ",
              "never a wrong lock, never a failed acquisition.")], size=10.5, gap=0)
_, tf = tb(s, 10.05, 6.06, 2.95, 0.80, anchor=MSO_ANCHOR.MIDDLE)
bullets(tf, [("Reproducible:  ",
              "every number regenerates from one command.")], size=10.5, gap=0)

# ================= SLIDE 5 — IMPACT =================
s = S[4]
clean_title(s); team_oval(s)
style_pointer_box(s, h=0.26, size=9)

_, tf = tb(s, 0.30, 1.64, 4.25, 0.42)
set_runs(tf.paragraphs[0], [("ECONOMIC FEASIBILITY", {"bold": True, "size": 13.5, "color": GREEN}),
                            ("  — the arithmetic that removes the cost barrier",
                             {"size": 10.5, "color": GRAY})])
pic(s, "c_econ_capital.jpg", 0.30, 2.06, w=4.06, border=RGBColor(0xD5, 0xDD, 0xE5))  # h 2.11
pic(s, "c_econ_stats.jpg", 0.30, 4.26, w=4.06, border=RGBColor(0xD5, 0xDD, 0xE5))    # h 1.30

pic(s, "c5_cards.jpg", 4.66, 1.62, w=8.20)                        # h = 3.96
pic(s, "c5_audience.jpg", 0.35, 5.66, w=5.72)                     # h = 1.25
pic(s, "c5_linkbox.png", 7.20, 5.66, w=5.75)                      # h = 1.26

# ================= SLIDE 6 — REFERENCES =================
s = S[5]
clean_title(s); team_oval(s)
style_pointer_box(s, h=0.22, size=9)


def scholar(t):
    return "https://scholar.google.com/scholar?q=" + quote_plus('"' + t + '"')


_, tf = tb(s, 0.42, 1.58, 6.0, 0.34)
set_runs(tf.paragraphs[0], [("Research foundations", {"bold": True, "size": 14, "color": NAVY})])
REFS = [
    ("Kaymak et al.,  ", "“A Survey on Acquisition, Tracking and Pointing Mechanisms for Mobile FSO,” IEEE Comm. Surveys & Tutorials, 2018.",
     scholar("A Survey on Acquisition Tracking and Pointing Mechanisms for Mobile Free-Space Optical Communications")),
    ("Kaushal & Kaddoum,  ", "“Optical Communication in Space: Challenges and Mitigation Techniques,” IEEE CS&T, 2017.",
     scholar("Optical Communication in Space Challenges and Mitigation Techniques")),
    ("Blackman & Popoli,  ", "Design and Analysis of Modern Tracking Systems — IMM filtering, CFAR detection.",
     scholar("Design and Analysis of Modern Tracking Systems Blackman Popoli")),
    ("Bar-Shalom et al.,  ", "Estimation with Applications to Tracking and Navigation — Kalman/IMM theory.",
     scholar("Estimation with Applications to Tracking and Navigation Bar-Shalom")),
    ("Smith, O.J.M.,  ", "“Closer Control of Loops with Dead Time” — the Smith predictor, 1957.",
     scholar("Closer Control of Loops with Dead Time Smith")),
    ("Vallado & Crawford,  ", "“SGP4 Orbit Determination” — AIAA 2008; real ISS TLE propagation.",
     scholar("SGP4 Orbit Determination Vallado Crawford")),
    ("Andrews & Phillips,  ", "Laser Beam Propagation through Random Media — turbulence & scintillation.",
     scholar("Laser Beam Propagation through Random Media Andrews Phillips")),
]
_, tf = tb(s, 0.42, 1.98, 6.15, 3.45)
first = True
for lead, rest, url in REFS:
    pr = tf.paragraphs[0] if first else tf.add_paragraph()
    first = False
    pr.space_after = Pt(4)
    set_runs(pr, [(lead, {"bold": True, "size": 10, "color": NAVY}),
                  (rest + "  ", {"size": 10, "color": NAVY})])
    lk = pr.add_run(); lk.text = "[link]"
    lk.font.name = "Arial"; lk.font.size = Pt(9.5)
    lk.font.color.rgb = CYAN; lk.font.underline = True
    lk.hyperlink.address = url
pic(s, "c6_logos.jpg", 1.10, 5.52, w=4.40)                        # h = 1.38

_, tf = tb(s, 6.95, 1.58, 6.0, 0.34)
set_runs(tf.paragraphs[0], [("Our work — open and verifiable", {"bold": True, "size": 14, "color": NAVY})])
_, tf = tb(s, 6.95, 1.98, 6.0, 1.42)
pr = tf.paragraphs[0]
set_runs(pr, [("Repository:  ", {"bold": True, "size": 11, "color": NAVY})])
lk = pr.add_run(); lk.text = "github.com/pranshukesavmishra/SIH-2026"
lk.font.name = "Arial"; lk.font.size = Pt(11)
lk.font.color.rgb = CYAN; lk.font.underline = True
lk.hyperlink.address = "https://github.com/pranshukesavmishra/SIH-2026"
pr.space_after = Pt(4)
pr2 = tf.add_paragraph()
set_runs(pr2, [("Live demo (replay console):  ", {"bold": True, "size": 11, "color": NAVY})])
lk = pr2.add_run(); lk.text = "zerodrift-fsoc-pat.netlify.app"
lk.font.name = "Arial"; lk.font.size = Pt(11)
lk.font.color.rgb = CYAN; lk.font.underline = True
lk.hyperlink.address = "https://zerodrift-fsoc-pat.netlify.app/"
pr2.space_after = Pt(4)
bullets(tf, [
    ("Inside:  ", "5,500+ documented lines · 83 automated tests · 6 validated scenarios · Monte-Carlo logs · technical report · user manual · demo video."),
    ("Reproducible:  ", "every figure regenerates from one command; every random draw is logged and replayable."),
], size=10.5, gap=4)

_, tf = tb(s, 6.95, 3.52, 5.5, 0.26)
set_runs(tf.paragraphs[0], [("Compared with existing approaches", {"bold": True, "size": 11.5, "color": NAVY})])
pic(s, "c6_table.jpg", 6.95, 3.80, w=5.30)                        # h = 1.86

shp, tf = card(s, 6.95, 5.74, 6.0, 1.16, fill=NAVY)
set_runs(tf.paragraphs[0], [("Why ZeroDrift wins on this problem",
                             {"bold": True, "size": 13, "color": WHITE})])
for line in ["All six mandatory deliverables already exist today",
             "Judged numbers, honestly measured against hidden ground truth",
             "AI where it earns its place — benchmarked against classical methods"]:
    p = tf.add_paragraph()
    set_runs(p, [("✓  ", {"bold": True, "size": 10.5, "color": RGBColor(0x7F, 0xD4, 0xA8)}),
                 (line, {"size": 10.5, "color": WHITE})])
    p.space_before = Pt(2)

# ============ delete the template's instructions slide ============
xml_slides = prs.slides._sldIdLst
xml_slides.remove(list(xml_slides)[6])

for _sl in prs.slides:
    for _sh in _sl.shapes:
        if _sh.has_text_frame and "@SIH Idea submission" in _sh.text_frame.text:
            for _par in _sh.text_frame.paragraphs:
                for _r in _par.runs:
                    if "@SIH Idea submission" in _r.text:
                        _r.text = "Team ZeroDrift · SIH 2026 · PS 26169"

core = prs.core_properties
core.title = "ZeroDrift — SIH 2026 Idea Submission (PS 26169, ISRO)"
core.author = "Team ZeroDrift, Jabalpur Engineering College"
core.subject = "AI-Based Virtual Camera Tracking for Coarse Alignment of Mobile FSOC Terminals"
core.keywords = "SIH 2026, PS 26169, ISRO, FSOC, ZeroDrift"
prs.save("ZeroDrift_SIH26169.pptx")
print("saved ZeroDrift_SIH26169.pptx")
